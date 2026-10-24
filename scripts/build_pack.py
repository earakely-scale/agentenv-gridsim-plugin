"""Build the GridSimEnv v1 packs from cache/raw: python scripts/build_pack.py [--workers N]."""

from __future__ import annotations

import argparse
import collections
import datetime as dt
import gzip
import hashlib
import io
import json
import math
import multiprocessing as mp
import sys
import time
import zipfile
from pathlib import Path

import numpy as np
import openpyxl

from gridsim_core import CONTRACT_VERSION
from gridsim_core.contract import era, price_violations
from gridsim_core.dp import solve_dp
from gridsim_core.grade import references
from gridsim_core.task import REGIONS, make_battery, point_kind, stable_seed, task_id

ROOT = Path(__file__).resolve().parents[1]
POINTS = ["HB_HOUSTON", "HB_NORTH", "HB_PAN", "HB_SOUTH", "HB_WEST", "LZ_AEN", "LZ_CPS", "LZ_HOUSTON", "LZ_LCRA",
          "LZ_NORTH", "LZ_RAYBN", "LZ_SOUTH", "LZ_WEST", "HB_HUBAVG", "HB_BUSAVG"]
SHOWCASE = {"uri": ("2021-02-10", "2021-02-22"), "elliott": ("2022-12-22", "2022-12-24"),
            "aug2023": ("2023-08-01", "2023-08-31"), "sep6_2023": ("2023-09-06", "2023-09-06"),
            "heather": ("2024-01-13", "2024-01-17")}
EVAL = ("2025-01-01", "2025-11-30")
RTCB = ("2025-12-05", "2026-10-17")
EP_MIX = {1: 0.40, 2: 0.45, 4: 0.15}
LOW_SPREAD_PER_MW = 22.0
GREEDY_EASY, GREEDY_EASY_DROP = 0.97, 0.80
TARGETS = {"train": 20000, "train_blind": 6000, "eval": 500, "eval_blind": 300, "eval_iid": 500, "rtcb": 250}
LETTERS = "abcdefghijklmnopqrstuvwxyz"
BINDING_SHARE = 0.10
TUNE_THRESHOLDS = (8.0, 15.0, 20.0, 22.0, 25.0, 30.0)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--raw", default=str(ROOT / "cache" / "raw"))
    ap.add_argument("--out", default=str(ROOT / "data"))
    ap.add_argument("--workers", type=int, default=max(1, (mp.cpu_count() or 2) - 2))
    args = ap.parse_args()
    t0 = time.time()
    raw, out = Path(args.raw), Path(args.out)
    sources = {p.name: sha256(p) for p in sorted((raw / "ercot").glob("DAMLZHBSPP_*.zip"))}
    eia_zips = sorted((raw / "eia860").glob("eia860*.zip"))
    sources.update({p.name: sha256(p) for p in eia_zips})
    dam, parse_stats = load_dam(raw / "ercot", ROOT / "cache" / "parsed", args.workers)
    log(f"parsed {len(dam)} point-days in {time.time() - t0:.0f}s")
    fit = eia_fit(eia_zips)
    table = letter_table(fit)
    kept, dedup_stats = dedup(dam)
    log(f"dedup kept {len(kept)} of {len(dam)}")
    split_of = {k: split(k[0]) for k in kept}
    ridge = fit_ridge(dam, [k for k in kept if split_of[k] == "train"])
    out.mkdir(parents=True, exist_ok=True)
    (out / "ridge-v1.json").write_text(json.dumps(ridge, indent=1) + "\n")
    era_flag = set(parse_stats["era_violations"])
    keys = [k for k in kept if split_of[k] is not None and (split_of[k] == "showcase" or f"{k[0]} {k[1]}" not in era_flag)]
    cands = [make_candidate(k, dam, table, salt) for k in keys for salt in (("", "2") if split_of[k] == "train" else ("",))]
    cands = list({c["task_id"]: c for c in cands}.values())
    log(f"{len(cands)} candidates; computing references on {args.workers} workers")
    with mp.Pool(args.workers, initializer=_init_worker, initargs=(ridge,)) as pool:
        cands = pool.map(_refs_worker, cands, chunksize=64)
    log(f"references done at {time.time() - t0:.0f}s")
    for c in cands:
        c["_split"] = split_of[(c["date"], c["settlement_point"])]
    packs, filt = select(cands)
    packs["fixtures"] = fixtures(packs, cands)
    cmd = "python scripts/build_pack.py " + " ".join(sys.argv[1:])
    for name, tasks in packs.items():
        write_pack(out, name, tasks, {"sources": sources, "filters": filt, "letters": table, "eia_fit": fit,
                                      "build_command": cmd.strip(), "parse": parse_stats, "dedup": dedup_stats})
    write_stats(out, packs, cands, filt, fit, parse_stats, dedup_stats, table)
    log(f"done in {time.time() - t0:.0f}s")


def log(msg: str):
    print(f"[build_pack] {msg}", flush=True)


def sha256(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


def dst_dates(year: int) -> set[str]:
    mar = dt.date(year, 3, 1)
    nov = dt.date(year, 11, 1)
    spring = mar + dt.timedelta(days=(6 - mar.weekday()) % 7 + 7)
    fall = nov + dt.timedelta(days=(6 - nov.weekday()) % 7)
    return {spring.isoformat(), fall.isoformat()}


def _parse_year(path: str) -> tuple[dict, dict]:
    z = zipfile.ZipFile(path)
    wb = openpyxl.load_workbook(io.BytesIO(z.read(z.namelist()[0])), read_only=True, data_only=True)
    vals, flags = collections.defaultdict(dict), collections.defaultdict(int)
    for ws in wb.worksheets:
        ws.reset_dimensions()
        rows = ws.iter_rows(values_only=True)
        head = [str(h).strip() for h in next(rows)]
        i_d, i_h, i_r = head.index("Delivery Date"), head.index("Hour Ending"), head.index("Repeated Hour Flag")
        i_p, i_v = head.index("Settlement Point"), head.index("Settlement Point Price")
        for r in rows:
            if r[i_d] is None:
                continue
            d = r[i_d]
            d = d.date().isoformat() if isinstance(d, dt.datetime) else dt.datetime.strptime(str(d), "%m/%d/%Y").date(
            ).isoformat()
            he = int(str(r[i_h]).split(":")[0])
            key = (d, str(r[i_p]).strip())
            if str(r[i_r]).strip().upper() == "Y":
                flags[key] += 1
                continue
            if he in vals[key]:
                flags[key] += 1
            vals[key][he] = float(r[i_v])
    return dict(vals), dict(flags)


def load_dam(ercot: Path, cache: Path, workers: int) -> tuple[dict, dict]:
    """{(date, point): np.array[24]} for clean days, plus parse statistics; cached by zip hashes."""
    zips = sorted(ercot.glob("DAMLZHBSPP_*.zip"))
    key = hashlib.sha256("".join(sha256(z) for z in zips).encode()).hexdigest()[:16]
    cache.mkdir(parents=True, exist_ok=True)
    cp = cache / f"dam-{key}.json.gz"
    if cp.exists():
        with gzip.open(cp, "rt") as f:
            blob = json.load(f)
        return {tuple(k.split("|")): np.array(v) for k, v in blob["dam"].items()}, blob["stats"]
    with mp.Pool(min(workers, len(zips))) as pool:
        parts = pool.map(_parse_year, [str(z) for z in zips])
    dam, st = {}, collections.Counter()
    dst = set().union(*(dst_dates(y) for y in range(2014, 2027)))
    points, era_bad = set(), []
    for vals, flags in parts:
        for k, hv in vals.items():
            points.add(k[1])
            st["point_days_raw"] += 1
            if k[1] not in POINTS:
                st["other_point"] += 1
                continue
            if k[0] in dst:
                st["dst_dropped"] += 1
                continue
            if flags.get(k) or sorted(hv) != list(range(1, 25)):
                st["not_24_dropped"] += 1
                continue
            p = np.array([hv[h] for h in range(1, 25)])
            if price_violations(k[0], p):
                era_bad.append(f"{k[0]} {k[1]}")
                st["era_flagged"] += 1
            dam[k] = p
            st["negative_hours"] += int((p < 0).sum())
            st["days_with_negative"] += int((p < 0).any())
            st["days_max_le_0"] += int(p.max() <= 0)
    stats = dict(st) | {"points_seen": sorted(points), "era_violations": sorted(era_bad), "clean_point_days": len(dam),
                        "first_date": {p: min(d for d, q in dam if q == p) for p in POINTS if any(q == p for _, q in dam)},
                        "last_date": max(d for d, _ in dam)}
    with gzip.open(cp, "wt") as f:
        json.dump({"dam": {f"{k[0]}|{k[1]}": v.tolist() for k, v in dam.items()}, "stats": stats}, f)
    return dam, stats


def dedup(dam: dict) -> tuple[list, dict]:
    """Within a date, drop a point whose prices match a kept one (corr > 0.999 and max abs diff < $1)."""
    by_date = collections.defaultdict(list)
    for d, p in dam:
        by_date[d].append(p)
    kept, dropped = [], collections.Counter()
    for d in sorted(by_date):
        chosen = []
        for p in sorted(by_date[d], key=POINTS.index):
            x = dam[(d, p)]
            dup = None
            for q in chosen:
                y = dam[(d, q)]
                if np.abs(x - y).max() < 1.0 and (x.std() < 1e-9 or y.std() < 1e-9 or np.corrcoef(x, y)[0, 1] > 0.999):
                    dup = q
                    break
            if dup:
                dropped[f"{p}->{dup}"] += 1
            else:
                chosen.append(p)
                kept.append((d, p))
    by_point = collections.Counter(k.split("->")[0] for k in dropped.elements())
    return kept, {"dropped": sum(dropped.values()), "by_point": dict(by_point), "pairs": dict(dropped.most_common(20))}


def in_range(d: str, r: tuple[str, str]) -> bool:
    return r[0] <= d <= r[1]


def split(d: str) -> str | None:
    """train | eval_iid | eval | rtcb | showcase | None (unused)."""
    if any(in_range(d, r) for r in SHOWCASE.values()):
        return "showcase"
    if d >= "2025-01-01":
        return "eval" if in_range(d, EVAL) else "rtcb" if in_range(d, RTCB) else None
    if (dt.date.fromisoformat(d) - dt.date(2014, 1, 1)).days % 10 == 9:
        return "eval_iid"
    return "train"


def lags(dam: dict, d: str, p: str) -> dict | None:
    day = dt.date.fromisoformat(d)
    prev = [dam.get(((day - dt.timedelta(days=i)).isoformat(), p)) for i in range(7, 0, -1)]
    if any(x is None for x in prev):
        return None
    return {"dam_prev_day": [round(float(v), 2) for v in prev[-1]],
            "dam_prev_7d": [[round(float(v), 2) for v in x] for x in prev]}


def fit_ridge(dam: dict, train_keys: list) -> dict:
    from gridsim_core.baselines import ridge_fit

    rows = []
    for d, p in train_keys:
        lg = lags(dam, d, p)
        if lg:
            rows.append((d, lg["dam_prev_day"], lg["dam_prev_7d"], dam[(d, p)]))
    best = None
    cut = "2023-01-01"
    fit_rows, val_rows = [r for r in rows if r[0] < cut], [r for r in rows if r[0] >= cut]
    from gridsim_core.baselines import ridge_features

    xv = np.stack([ridge_features(r[0], r[1], r[2]) for r in val_rows])
    yv = np.asarray([r[3] for r in val_rows])
    scores = {}
    for lam in (1e2, 1e3, 1e4, 1e5, 1e6, 1e7):
        m = ridge_fit(fit_rows, lam)
        pred = np.einsum("nhf,hf->nh", xv, np.asarray(m["coef"]))
        scores[lam] = float(np.mean(np.abs(pred - yv)))
        if best is None or scores[lam] < scores[best]:
            best = lam
    model = ridge_fit(rows, best)
    model["fit_dates"] = "train split 2014-2024 (excludes every-10th eval-iid dates and showcase dates)"
    model["lambda_selection"] = {"fit_before": cut, "val_mae_by_lambda": scores}
    pers = float(np.mean(np.abs(np.asarray([r[1] for r in val_rows]) - yv)))
    model["val_mae_persistence"] = pers
    log(f"ridge lambda={best} val MAE {scores[best]:.2f} vs persistence {pers:.2f} on {len(val_rows)} rows")
    return model


def eia_fit(zips: list[Path]) -> dict:
    """ERCO operable battery (P, E/P) distribution from EIA-860 Schedule 3_4 joined to the plant BA code."""
    z = zipfile.ZipFile(zips[-1])
    names = z.namelist()
    sto = next(n for n in names if n.startswith("3_4_Energy_Storage"))
    plant = next(n for n in names if n.startswith("2___Plant"))
    wb = openpyxl.load_workbook(io.BytesIO(z.read(plant)), read_only=True, data_only=True)
    rows = wb.worksheets[0].iter_rows(values_only=True)
    next(rows)
    head = list(next(rows))
    ba = {r[head.index("Plant Code")]: r[head.index("Balancing Authority Code")] for r in rows}
    wb = openpyxl.load_workbook(io.BytesIO(z.read(sto)), read_only=True, data_only=True)
    rows = wb["Operable"].iter_rows(values_only=True)
    next(rows)
    head = list(next(rows))
    units = []
    for r in rows:
        if ba.get(r[head.index("Plant Code")]) != "ERCO" or r[head.index("Technology")] != "Batteries":
            continue
        p = _num(r[head.index("Maximum Discharge Rate (MW)")]) or _num(r[head.index("Nameplate Capacity (MW)")])
        e = _num(r[head.index("Nameplate Energy Capacity (MWh)")])
        if p and e and p >= 1:
            units.append((p, e / p, int(_num(r[head.index("Operating Year")]) or 0)))
    pw = np.array([u[0] for u in units])
    ep = np.array([u[1] for u in units])
    bucket = np.where(ep < 1.5, 1, np.where(ep < 3.0, 2, 4))
    by_count = {f"{b}h": round(float((bucket == b).mean()), 4) for b in (1, 2, 4)}
    by_mw = {f"{b}h": round(float(pw[bucket == b].sum() / pw.sum()), 4) for b in (1, 2, 4)}
    recent = np.array([u[2] >= 2023 for u in units])
    return {"source": zips[-1].name, "schedule": sto, "ba": "ERCO", "units": len(units),
            "total_mw": round(float(pw.sum()), 1), "mw_weighted_mean_duration_h": round(float((pw * ep).sum() / pw.sum()), 3),
            "p_mw_quantiles": {q: round(float(np.quantile(pw, q / 100)), 1) for q in (10, 25, 50, 75, 90)},
            "p_mw_quantiles_mw_weighted": {q: round(_wq(pw, pw, q / 100), 1) for q in (10, 25, 50, 75, 90)},
            "ep_h_quantiles_mw_weighted": {q: round(_wq(ep, pw, q / 100), 2) for q in (10, 25, 50, 75, 90)},
            "ep_h_quantiles": {q: round(float(np.quantile(ep, q / 100)), 2) for q in (10, 25, 50, 75, 90)},
            "ep_bucket_share_by_count": by_count, "ep_bucket_share_by_mw": by_mw,
            "ep_bucket_share_by_mw_since_2023": {
                f"{b}h": round(float(pw[(bucket == b) & recent].sum() / max(pw[recent].sum(), 1)), 4) for b in (1, 2, 4)},
            "buckets": "1h: E/P < 1.5 h, 2h: 1.5-3 h, 4h: >= 3 h",
            "archetype_mix_used": {f"{k}h": v for k, v in EP_MIX.items()}}


def _wq(x: np.ndarray, w: np.ndarray, q: float) -> float:
    o = np.argsort(x)
    c = np.cumsum(w[o]) / w.sum()
    return float(x[o][np.searchsorted(c, q)])


def _num(x) -> float | None:
    try:
        return float(x)
    except (TypeError, ValueError):
        return None


def letter_table(fit: dict) -> dict:
    """26 battery letters per E/P class, a deterministic balanced design over the archetype factors."""
    rng = np.random.default_rng(4329)
    qs = fit["p_mw_quantiles_mw_weighted"]
    p_levels = sorted({_nice(fit["p_mw_quantiles"][50]), _nice(qs[25]), _nice(qs[50]), _nice(qs[75])})
    table = {}
    for ep in EP_MIX:
        rows = []
        etas = rng.permutation(np.resize([0.80, 0.85, 0.90], 26))
        degs = rng.permutation(np.resize([0.0, 4.0, 8.0], 26))
        cyc = rng.permutation(np.resize([1.0, 1.5, 2.0], 26))
        ks = rng.permutation(np.resize([2, 3, 4], 26))
        ps = rng.permutation(np.resize(p_levels, 26))
        asym = rng.permutation([True] * 6 + [False] * 20)
        above = rng.permutation([True] * 4 + [False] * 22)
        s0s = rng.permutation(np.resize(["0", "half", "full"], 26))
        for i, letter in enumerate(LETTERS):
            k = int(ks[i])
            n = ep * k
            while n > 16:
                k -= 1
                n = ep * k
            k_ch = k
            if asym[i]:
                k_ch = k - 1 if k > 2 else k + 1
            s0r = str(s0s[i])
            if above[i]:
                s0r = "0" if s0r == "full" else s0r
            s0 = {"0": 0, "half": n // 2, "full": n}[s0r]
            s_end = (n // 2 if s0 == 0 else n) if above[i] else (s0 if i % 2 == 0 else min(s0, n // 2) if i % 4 == 1 else 0)
            q = _nice_q(float(ps[i]) / k)
            p = q * k
            rows.append({"letter": letter, "ep_h": ep, "p_dis_mw": p, "k_dis": k, "k_ch": k_ch, "unit_mwh": q,
                         "soc_levels": n, "eta": float(etas[i]), "c_deg": float(degs[i]), "cycles": float(cyc[i]),
                         "s0": s0, "s_end_min": s_end})
        table[f"{ep}h"] = rows
    return table


def _nice_q(x: float) -> float:
    grid = (2.5, 5.0, 7.5, 10.0, 12.5, 15.0, 20.0, 25.0, 30.0, 40.0, 50.0, 75.0)
    return min(grid, key=lambda g: abs(math.log(g / x)))


def _nice(x: float) -> float:
    for v in (10, 20, 50, 100, 150, 200, 300):
        if x <= v * 1.2:
            return float(v)
    return 300.0


def battery_for(row: dict) -> dict:
    return make_battery(row["p_dis_mw"], row["p_dis_mw"] * row["ep_h"], row["k_dis"], k_ch=row["k_ch"],
                        eta=row["eta"], c_deg=row["c_deg"], cycles=row["cycles"], s0=row["s0"],
                        s_end_min=row["s_end_min"], archetype=f"{row['ep_h']}h-{row['letter']}")


def pick_archetype(d: str, p: str, table: dict, salt: str = "") -> dict:
    h = stable_seed(f"arch{salt}|{d}|{p}")
    u = (h % 100000) / 100000
    acc, ep = 0.0, 4
    for e, w in EP_MIX.items():
        acc += w
        if u < acc:
            ep = e
            break
    return table[f"{ep}h"][(h // 100000) % 26]


def make_candidate(key: tuple, dam: dict, table: dict, salt: str = "") -> dict:
    d, p = key
    row = pick_archetype(d, p, table, salt)
    b = battery_for(row)
    t = {"task_id": task_id(d, p, b["archetype"]), "contract_version": CONTRACT_VERSION, "pack": None, "split": None,
         "date": d, "settlement_point": p, "point_kind": point_kind(p), "region": REGIONS[p],
         "prices": [round(float(v), 2) for v in dam[key]], "blind": lags(dam, d, p), "battery": b}
    return t


_RIDGE = None


def _init_worker(ridge):
    global _RIDGE
    _RIDGE = ridge


def _refs_worker(t: dict) -> dict:
    t["refs"] = references(t, ridge=_RIDGE if t["blind"] else None)
    f = t["refs"]["foresight"]
    t["difficulty"] = round(1.0 - f["baselines"]["greedy"]["reward"], 6)
    t["tags"] = tags(t)
    t["curriculum"] = curriculum(t)
    return t


def tags(t: dict) -> dict:
    p = np.asarray(t["prices"])
    b = t["battery"]
    f = t["refs"]["foresight"]
    p_opt, margin = f["p_opt"], BINDING_SHARE * max(f["p_opt"] - f["p_idle"], 0.0)
    binding = "none"
    for name, change in (("cycle_cap", {"cycle_cap_units": 24 * b["k_dis"]}),
                         ("end_soc", {"s_end_min": min(b["s0"], b["s_end_min"])}),
                         ("rate", {"k_dis": b["soc_levels"], "k_ch": b["soc_levels"]})):
        relaxed = dict(b, **change)
        if relaxed != b and solve_dp({"prices": t["prices"], "battery": relaxed})[0] > p_opt + margin + 1e-6:
            binding = name
            break
    m = int(t["date"][5:7])
    season = "winter" if m in (12, 1, 2) else "spring" if m <= 5 else "summer" if m <= 8 else "fall"
    return {"neg_hours": int((p < 0).sum()), "spike": bool(p.max() >= max(200.0, 3 * float(np.median(p)))),
            "season": season, "binding": binding, "era": era(t["date"])["name"]}


def curriculum(t: dict) -> str:
    g = t["tags"]
    if g["binding"] in ("cycle_cap", "end_soc") or g["neg_hours"] > 0:
        return "C2"
    return "C0" if t["difficulty"] <= 0.15 else "C1"


def low_spread(t: dict, thr: float = LOW_SPREAD_PER_MW) -> bool:
    f = t["refs"]["foresight"]
    return f["p_opt"] - f["p_idle"] < thr * t["battery"]["p_dis_mw"]


def blind_ok(t: dict, thr: float = LOW_SPREAD_PER_MW) -> bool:
    bl = t["refs"]["blind"]
    return bool(bl) and bl["p_anchor"] - t["refs"]["foresight"]["p_idle"] >= thr * t["battery"]["p_dis_mw"]


def greedy_easy_drop(t: dict) -> bool:
    if t["refs"]["foresight"]["baselines"]["greedy"]["reward"] < GREEDY_EASY:
        return False
    return (stable_seed("easy|" + t["task_id"]) % 1000) / 1000 < GREEDY_EASY_DROP


def tuning_table(cands: list) -> list[dict]:
    rows = []
    ev = [c for c in cands if c["_split"] in ("eval", "train")]
    for thr in TUNE_THRESHOLDS:
        for drop in (0.0, 0.5, 0.8):
            for sp in ("train", "eval"):
                pool = [c for c in ev if c["_split"] == sp]
                ls = [c for c in pool if not low_spread(c, thr)]
                kept = [c for c in ls if not (c["refs"]["foresight"]["baselines"]["greedy"]["reward"] >= GREEDY_EASY
                                              and (stable_seed("easy|" + c["task_id"]) % 1000) / 1000 < drop)]
                g = [c["refs"]["foresight"]["baselines"]["greedy"]["reward"] for c in kept]
                rows.append({"split": sp, "low_spread_per_mw": thr, "greedy_easy_drop": drop,
                             "low_spread_share": round(1 - len(ls) / max(len(pool), 1), 4),
                             "kept": len(kept), "greedy_mean": round(float(np.mean(g)), 4) if g else None})
    return rows


def stratified(tasks: list, n: int, key, salt: str) -> list:
    """Round-robin over strata (in hash order inside each stratum) until n tasks."""
    strata = collections.defaultdict(list)
    for t in sorted(tasks, key=lambda t: stable_seed(salt + t["task_id"])):
        strata[key(t)].append(t)
    order = sorted(strata, key=lambda k: stable_seed(salt + repr(k)))
    out, i = [], 0
    while len(out) < n and any(strata.values()):
        k = order[i % len(order)]
        if strata[k]:
            out.append(strata[k].pop(0))
        i += 1
    return out


def select(cands: list) -> tuple[dict, dict]:
    by = collections.defaultdict(list)
    for c in cands:
        by[c["_split"]].append(c)
    filt = {"low_spread_per_mw": LOW_SPREAD_PER_MW, "greedy_easy": GREEDY_EASY, "greedy_easy_drop": GREEDY_EASY_DROP,
            "dedup": "corr > 0.999 and max abs diff < $1 within a date", "per_split": {}, "tuning": tuning_table(cands)}
    packs = {}
    for sp in ("train", "eval", "eval_iid", "rtcb"):
        pool = by[sp]
        ls = [c for c in pool if low_spread(c)]
        kept = [c for c in pool if not low_spread(c)]
        easy = [c for c in kept if greedy_easy_drop(c)]
        kept = [c for c in kept if not greedy_easy_drop(c)]
        bl = [c for c in kept if c["blind"]]
        bl_ok = [c for c in bl if blind_ok(c)]
        filt["per_split"][sp] = {"candidates": len(pool), "low_spread_dropped": len(ls),
                                 "low_spread_share": round(len(ls) / max(len(pool), 1), 4),
                                 "greedy_easy_dropped": len(easy), "kept": len(kept),
                                 "blind_candidates": len(bl), "blind_anchor_dropped": len(bl) - len(bl_ok),
                                 "blind_anchor_dropped_share": round(1 - len(bl_ok) / max(len(bl), 1), 4)}
        if sp == "train":
            sel = sorted(kept, key=lambda t: stable_seed("train|" + t["task_id"]))[:TARGETS["train"]]
            blind_ids = {t["task_id"] for t in sorted([t for t in sel if t["blind"] and blind_ok(t)],
                                                      key=lambda t: stable_seed("tb|" + t["task_id"]))[:TARGETS["train_blind"]]}
        else:
            key = (lambda t: (t["settlement_point"], t["battery"]["archetype"].split("-")[0], t["date"][5:7]))
            sel = stratified(kept, TARGETS[sp], key, sp + "|")
            blind_ids = set()
            if sp == "eval":
                blind_ids = {t["task_id"] for t in stratified([t for t in sel if t["blind"] and blind_ok(t)],
                                                              TARGETS["eval_blind"], key, "eb|")}
            elif sp == "rtcb":
                blind_ids = {t["task_id"] for t in sel if t["blind"] and blind_ok(t)}
        packs[sp] = [finalize(t, sp, t["task_id"] in blind_ids) for t in sel]
    packs["showcase"] = [finalize(t, "showcase", bool(t["blind"]) and blind_ok(t), keep_flags=True)
                         for t in sorted(by["showcase"], key=lambda t: (t["date"], POINTS.index(t["settlement_point"])))]
    return packs, filt


PACK_NAMES = {"train": "ercot-dam-v1-train", "eval": "ercot-dam-v1-eval", "eval_iid": "ercot-dam-v1-eval-iid",
              "rtcb": "ercot-dam-v1-rtcb", "showcase": "showcase", "fixtures": "fixtures"}
SPLIT_NAMES = {"train": "train", "eval": "eval_2025", "eval_iid": "eval_iid", "rtcb": "rtcb", "showcase": "showcase"}


def finalize(t: dict, sp: str, blind: bool, keep_flags: bool = False) -> dict:
    out = {k: v for k, v in t.items() if not k.startswith("_")}
    out["pack"], out["split"] = PACK_NAMES[sp], SPLIT_NAMES[sp]
    if not blind:
        out["blind"] = None
        out["refs"] = dict(out["refs"], blind=None)
    out["tags"] = dict(out["tags"], blind_eval=blind)
    if keep_flags:
        out["tags"]["low_spread"] = low_spread(t)
    order = ["task_id", "contract_version", "pack", "split", "date", "settlement_point", "point_kind", "region",
             "prices", "blind", "battery", "refs", "difficulty", "curriculum", "tags"]
    return {k: out[k] for k in order}


def fixtures(packs: dict, cands: list) -> list:
    """20 CI tasks: eval and eval-iid mix, every E/P class, blind ones, an s_end_min > s0 day and a max-price <= 0 day."""
    pool = packs["eval"] + packs["eval_iid"] + packs["rtcb"]
    out, seen = [], set()

    def take(pred, n):
        for t in sorted(pool, key=lambda t: stable_seed("fx|" + t["task_id"])):
            if len([x for x in out if pred(x)]) >= n or len(out) >= 20:
                return
            if pred(t) and t["task_id"] not in seen:
                out.append(t)
                seen.add(t["task_id"])

    take(lambda t: t["tags"]["blind_eval"], 4)
    take(lambda t: t["battery"]["s_end_min"] > t["battery"]["s0"], 2)
    for ep in ("1h", "2h", "4h"):
        take(lambda t, ep=ep: t["battery"]["archetype"].startswith(ep), 3)
    take(lambda t: t["tags"]["binding"] == "cycle_cap", 1)
    take(lambda t: t["tags"]["neg_hours"] > 0, 1)
    nonpos = sorted([c for c in cands if max(c["prices"]) <= 0], key=lambda t: stable_seed("np|" + t["task_id"]))
    if nonpos:
        out.append(finalize(nonpos[0], nonpos[0]["_split"], False))
    take(lambda t: True, 20)
    return [dict(t, pack="fixtures") for t in out[:20]]


def write_pack(out: Path, name: str, tasks: list, meta: dict):
    d = out / PACK_NAMES[name]
    d.mkdir(parents=True, exist_ok=True)
    with gzip.GzipFile(d / "tasks.jsonl.gz", "wb", mtime=0) as f:
        for t in tasks:
            f.write((json.dumps(t, separators=(",", ":")) + "\n").encode())
    ep = collections.Counter(t["battery"]["archetype"].split("-")[0] for t in tasks)
    manifest = {"pack": PACK_NAMES[name], "contract_version": CONTRACT_VERSION, "count": len(tasks),
                "count_blind": sum(1 for t in tasks if t["tags"]["blind_eval"]),
                "blind_task_ids": sorted(t["task_id"] for t in tasks if t["tags"]["blind_eval"])
                if name != "train" else None,
                "date_range": [min(t["date"] for t in tasks), max(t["date"] for t in tasks)] if tasks else None,
                "by_ep": dict(sorted(ep.items())),
                "by_point": dict(sorted(collections.Counter(t["settlement_point"] for t in tasks).items())),
                "by_curriculum": dict(sorted(collections.Counter(t["curriculum"] for t in tasks).items())),
                "means": baseline_means(tasks), "tasks_sha256": sha256(d / "tasks.jsonl.gz"),
                "holdouts": {"eval": EVAL, "rtcb": RTCB, "showcase": SHOWCASE,
                             "eval_iid": "every 10th calendar date from 2014-01-10 through 2024", "train": "2014-2024 rest"},
                "archetype_mix": {f"{k}h": v for k, v in EP_MIX.items()}, **meta}
    (d / "manifest.json").write_text(json.dumps(manifest, indent=1) + "\n")


def baseline_means(tasks: list) -> dict:
    if not tasks:
        return {}
    out = {}
    for name in ("clock", "threshold", "greedy", "random"):
        out[name] = round(float(np.mean([t["refs"]["foresight"]["baselines"][name]["reward"] for t in tasks])), 4)
    bl = [t["refs"]["blind"] for t in tasks if t["refs"]["blind"]]
    if bl:
        out["blind_anchor_r_opt"] = round(float(np.mean([b["r_opt_anchor"] for b in bl])), 4)
        out["blind_persist_r_opt"] = round(float(np.mean([_r_opt(t, t["refs"]["blind"]["p_persist"]) for t in tasks
                                                          if t["refs"]["blind"]])), 4)
    gaps = lp_gaps(tasks)
    if gaps.size:
        out["lp_dp_gap_median"] = round(float(np.median(gaps)), 5)
        out["lp_dp_gap_p95"] = round(float(np.quantile(gaps, 0.95)), 5)
    return out


def _r_opt(t: dict, p: float) -> float:
    from gridsim_core.grade import reward_value

    f = t["refs"]["foresight"]
    return reward_value(p, f["p_idle"], f["p_opt"])


def lp_gaps(tasks: list) -> np.ndarray:
    """(P_lp - P_opt) / (P_lp - P_idle): the relaxation's headroom over the integer optimum, on the reward scale."""
    out = []
    for t in tasks:
        f = t["refs"]["foresight"]
        if f["p_lp"] is not None and f["p_lp"] - f["p_idle"] > 1e-6:
            out.append((f["p_lp"] - f["p_opt"]) / (f["p_lp"] - f["p_idle"]))
    return np.asarray(out)


def write_stats(out: Path, packs: dict, cands: list, filt: dict, fit: dict, parse_stats: dict, dedup_stats: dict,
                table: dict):
    lines = ["# GridSimEnv v1 pack statistics", "",
             f"Built by `scripts/build_pack.py` (contract {CONTRACT_VERSION}). All numbers are measured on the packs.", ""]
    lines += ["## Packs", "", "| pack | tasks | blind | clock | threshold | greedy | random | LP-DP gap median | p95 |",
              "|---|---|---|---|---|---|---|---|---|"]
    for name, tasks in packs.items():
        m = baseline_means(tasks)
        lines.append(f"| {PACK_NAMES[name]} | {len(tasks)} | {sum(t['tags']['blind_eval'] for t in tasks)} | "
                     f"{m.get('clock')} | {m.get('threshold')} | {m.get('greedy')} | {m.get('random')} | "
                     f"{m.get('lp_dp_gap_median')} | {m.get('lp_dp_gap_p95')} |")
    lines += ["", ("Baseline columns are mean foresight rewards (idle is 0.2 by construction). The LP-DP gap is "
              "(P_lp - P_opt) / (P_lp - P_idle)."), ""]
    lines += ["## Blind tier", "", "| pack | blind tasks | anchor r_opt | persistence r_opt |", "|---|---|---|---|"]
    for name, tasks in packs.items():
        bt = [t for t in tasks if t["tags"]["blind_eval"]]
        if bt:
            m = baseline_means(bt)
            lines.append(f"| {PACK_NAMES[name]} | {len(bt)} | {m.get('blind_anchor_r_opt')} | "
                         f"{m.get('blind_persist_r_opt')} |")
    lines += ["", "## Filters", "", ("| split | candidates | low-spread dropped | share | greedy-easy dropped | kept | "
              "blind candidates | blind anchor dropped (share) |"), "|---|---|---|---|---|---|---|---|"]
    for sp, s in filt["per_split"].items():
        lines.append(f"| {sp} | {s['candidates']} | {s['low_spread_dropped']} | {s['low_spread_share']} | "
                     f"{s['greedy_easy_dropped']} | {s['kept']} | {s['blind_candidates']} | "
                     f"{s['blind_anchor_dropped']} ({s['blind_anchor_dropped_share']}) |")
    lines += ["", (f"Low-spread: drop when P_opt - P_idle < ${LOW_SPREAD_PER_MW:.0f} per MW of p_dis. Greedy-easy: drop "
              f"{GREEDY_EASY_DROP:.0%} of tasks with greedy >= {GREEDY_EASY} (hash-chosen). Blind: a task leaves the "
              "blind tier when P_anchor - P_idle is below the low-spread margin."), "",
              "### Tuning table", "", "| split | $/MW | greedy-easy drop | low-spread share | kept | greedy mean |",
              "|---|---|---|---|---|---|"]
    for r in filt["tuning"]:
        lines.append(f"| {r['split']} | {r['low_spread_per_mw']} | {r['greedy_easy_drop']} | {r['low_spread_share']} | "
                     f"{r['kept']} | {r['greedy_mean']} |")
    for title, keyf in (("point", lambda t: t["settlement_point"]), ("archetype class", lambda t: t["battery"][
            "archetype"].split("-")[0]), ("season", lambda t: t["tags"]["season"]), ("curriculum", lambda t: t[
            "curriculum"]), ("binding constraint", lambda t: t["tags"]["binding"])):
        names = [n for n in packs if n != "fixtures"]
        keys = sorted({keyf(t) for n in names for t in packs[n]})
        lines += ["", f"## Counts by {title}", "", "| " + title + " | " + " | ".join(PACK_NAMES[n] for n in names) + " |",
                  "|---" * (len(names) + 1) + "|"]
        for k in keys:
            lines.append(f"| {k} | " + " | ".join(str(sum(keyf(t) == k for t in packs[n])) for n in names) + " |")
    allc = cands
    neg = [c for c in allc if c["tags"]["neg_hours"] > 0]
    lines += ["", "## Prices", "",
              (f"- Clean point-days after parsing: {parse_stats['clean_point_days']} (DST dropped "
              f"{parse_stats.get('dst_dropped', 0)}, not-24 dropped {parse_stats.get('not_24_dropped', 0)}, above the era "
              f"cap/floor {parse_stats.get('era_flagged', 0)}, kept for showcase only: "
              f"{', '.join(parse_stats['era_violations'])})."),
              f"- Within-date dedup dropped {dedup_stats['dropped']} point-days: {dedup_stats['by_point']}.",
              (f"- Candidates with any negative hour: {len(neg)} of {len(allc)} ({len(neg) / max(len(allc), 1):.1%}); "
              f"negative hours: {sum(c['tags']['neg_hours'] for c in allc)} of {24 * len(allc)} "
              f"({sum(c['tags']['neg_hours'] for c in allc) / max(24 * len(allc), 1):.2%}); days with max price <= 0: "
              f"{sum(max(c['prices']) <= 0 for c in allc)}."),
              f"- First date per point: {parse_stats['first_date']}.", ""]
    gaps = lp_gaps(allc)
    if gaps.size:
        qs = {q: float(np.quantile(gaps, q / 100)) for q in (50, 75, 90, 95, 99)}
        lines += ["## LP-DP gap (all candidates)", "", "| quantile | " + " | ".join(f"p{q}" for q in qs) + " | max |",
                  "|---" * (len(qs) + 2) + "|", "| (P_lp - P_opt)/(P_lp - P_idle) | " + " | ".join(
                      f"{v:.4f}" for v in qs.values()) + f" | {gaps.max():.4f} |", ""]
    lines += ["## Battery archetypes (EIA-860 fit)", "", "```json", json.dumps(fit, indent=1), "```", "",
              ("The E/P mix used for tasks is the design mix (1 h 40%, 2 h 45%, 4 h 15%); the EIA ERCO fit above is the "
              "check. Letters per class (manifest `letters`):"), "",
              "| code | P MW | k_dis | k_ch | q MWh | N | eta | c_deg | cycles | s0 | s_end_min |",
              "|---|---|---|---|---|---|---|---|---|---|---|"]
    for ep, rows in table.items():
        for r in rows:
            lines.append(f"| {ep}-{r['letter']} | {r['p_dis_mw']:g} | {r['k_dis']} | {r['k_ch']} | {r['unit_mwh']:g} | "
                         f"{r['soc_levels']} | {r['eta']} | {r['c_deg']:g} | {r['cycles']} | {r['s0']} | "
                         f"{r['s_end_min']} |")
    (out / "stats.md").write_text("\n".join(lines) + "\n")


if __name__ == "__main__":
    main()
